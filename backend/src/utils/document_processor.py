"""Document processing utilities for converting various file formats to Markdown."""

import os
import base64
import logging
from typing import Optional, List
from docx import Document
import PyPDF2
import pdfplumber
import markdown2
import mammoth
from pathlib import Path
from datetime import datetime
from io import BytesIO

# Optional imports for OCR
try:
    from pdf2image import convert_from_path
    PDF2IMAGE_AVAILABLE = True
except ImportError:
    PDF2IMAGE_AVAILABLE = False

try:
    from openai import OpenAI
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


class DocumentProcessor:
    """Handles conversion of various document formats to Markdown."""
    
    @staticmethod
    def word_to_markdown(file_path: str) -> str:
        """
        Convert Word document to Markdown text using mammoth for better conversion.

        Args:
            file_path: Path to the Word document

        Returns:
            Markdown text content
        """
        # Check if it's an old .doc format (not supported)
        if file_path.lower().endswith('.doc') and not file_path.lower().endswith('.docx'):
            raise ValueError(
                "Old .doc format is not supported. Please convert to .docx format first. "
                "You can do this by opening the file in Microsoft Word and saving as .docx"
            )
        
        try:
            # Try mammoth first for better conversion
            with open(file_path, "rb") as docx_file:
                result = mammoth.convert_to_markdown(docx_file)
                markdown_content = result.value

                # Clean up the markdown
                lines = markdown_content.split('\n')
                cleaned_lines = []
                for line in lines:
                    stripped = line.strip()
                    if stripped:
                        cleaned_lines.append(stripped)
                        cleaned_lines.append("")  # Add blank line

                return "\n".join(cleaned_lines).strip()
        except Exception as e:
            # Fallback to python-docx if mammoth fails
            doc = Document(file_path)
            markdown_lines = []

            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue

                # Check paragraph style for headings
                if paragraph.style.name.startswith('Heading'):
                    level = paragraph.style.name.replace('Heading ', '')
                    try:
                        level_num = int(level)
                        markdown_lines.append(f"{'#' * level_num} {text}")
                    except ValueError:
                        markdown_lines.append(text)
                else:
                    markdown_lines.append(text)

                markdown_lines.append("")  # Add blank line between paragraphs

            return "\n".join(markdown_lines)
    
    @staticmethod
    def pdf_to_markdown(file_path: str) -> str:
        """
        Convert PDF to Markdown text using pdfplumber for better text extraction.
        Falls back to PyPDF2 if pdfplumber fails.
        
        Args:
            file_path: Path to the PDF file
            
        Returns:
            Markdown text content
        """
        import logging
        markdown_lines = []
        
        # Try pdfplumber first
        try:
            with pdfplumber.open(file_path) as pdf:
                logging.info(f"PDF has {len(pdf.pages)} pages")
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    logging.info(f"Page {i+1} extracted text length: {len(text) if text else 0}")
                    if text:
                        # Split into paragraphs (double newlines)
                        paragraphs = text.split('\n\n')
                        for para in paragraphs:
                            # Clean up single newlines within paragraphs
                            cleaned = para.replace('\n', ' ').strip()
                            if cleaned:
                                markdown_lines.append(cleaned)
                                markdown_lines.append("")  # Blank line between paragraphs
        except Exception as e:
            logging.error(f"pdfplumber failed: {e}")
        
        # If pdfplumber didn't extract anything, try PyPDF2
        if not markdown_lines:
            logging.info("Trying PyPDF2 fallback")
            try:
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    logging.info(f"PyPDF2: PDF has {len(reader.pages)} pages")
                    for i, page in enumerate(reader.pages):
                        text = page.extract_text()
                        logging.info(f"PyPDF2 Page {i+1} extracted text length: {len(text) if text else 0}")
                        if text:
                            # Clean up and add to markdown
                            lines = text.split('\n')
                            for line in lines:
                                cleaned = line.strip()
                                if cleaned:
                                    markdown_lines.append(cleaned)
                            markdown_lines.append("")  # Blank line between pages
            except Exception as e:
                logging.error(f"PyPDF2 also failed: {e}")
        
        result = "\n".join(markdown_lines)
        logging.info(f"Final markdown content length: {len(result)}")
        
        # If still empty, the PDF might be image-based (scanned) - try OCR
        if not result.strip():
            logging.info("Text extraction failed, attempting OCR with OpenAI Vision")
            ocr_result = DocumentProcessor.ocr_pdf_with_openai(file_path)
            if ocr_result:
                return ocr_result
            return "# PDF Text Extraction Failed\n\nThis PDF appears to be image-based (scanned) and OCR processing failed. Please check your OpenAI API key configuration."
        
        return result
    
    @staticmethod
    def image_to_markdown(file_path: str, ocr_text: str) -> str:
        """
        Convert OCR'd image text to Markdown.

        Args:
            file_path: Path to the image file
            ocr_text: OCR extracted text from AWS Textract or similar service

        Returns:
            Markdown text content
        """
        # For now, just format the OCR text as paragraphs
        # In production, this would use AWS Textract response structure
        markdown_lines = []

        paragraphs = ocr_text.split('\n\n')
        for para in paragraphs:
            cleaned = para.strip()
            if cleaned:
                markdown_lines.append(cleaned)
                markdown_lines.append("")

        return "\n".join(markdown_lines)

    @staticmethod
    def text_to_markdown(file_path: str) -> str:
        """
        Convert plain text file to Markdown.

        Args:
            file_path: Path to the text file

        Returns:
            Markdown text content
        """
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()

        return content
    
    @staticmethod
    def validate_markdown(content: str) -> bool:
        """
        Validate that the content is valid Markdown.

        Args:
            content: Markdown content to validate

        Returns:
            True if valid, False otherwise
        """
        # Simple validation - just check that content exists and is not empty
        # Any text is valid markdown, so we don't need strict validation
        if not content or not content.strip():
            return False

        return True
    
    @staticmethod
    def extract_identifiers(content: str) -> list[str]:
        """
        Extract all identifiers from Markdown content (e.g., <<identifier>>).
        
        Args:
            content: Markdown content
            
        Returns:
            List of unique identifiers
        """
        import re
        pattern = r'<<([^>]+)>>'
        matches = re.findall(pattern, content)
        return list(set(matches))
    
    @staticmethod
    def save_uploaded_file(file_content: bytes, filename: str, upload_dir: str) -> str:
        """
        Save uploaded file to storage.
        
        Args:
            file_content: File content as bytes
            filename: Original filename
            upload_dir: Directory to save the file
            
        Returns:
            Path to saved file
        """
        # Create upload directory if it doesn't exist
        Path(upload_dir).mkdir(parents=True, exist_ok=True)
        
        # Generate unique filename to avoid conflicts
        from datetime import datetime
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        name, ext = os.path.splitext(filename)
        unique_filename = f"{name}_{timestamp}{ext}"
        
        file_path = os.path.join(upload_dir, unique_filename)
        
        with open(file_path, 'wb') as f:
            f.write(file_content)
        
        return file_path
    
    @staticmethod
    def get_file_type(filename: str) -> Optional[str]:
        """
        Determine file type from filename extension.

        Args:
            filename: Name of the file

        Returns:
            File type: 'word', 'pdf', 'image', 'text', or None
        """
        ext = os.path.splitext(filename)[1].lower()

        if ext in ['.doc', '.docx']:
            return 'word'
        elif ext == '.pdf':
            return 'pdf'
        elif ext in ['.jpg', '.jpeg', '.png', '.tiff', '.tif', '.bmp']:
            return 'image'
        elif ext == '.txt':
            return 'text'

        return None

    @staticmethod
    def save_markdown_file(
        markdown_content: str,
        template_name: str,
        username: str,
        base_dir: str = "document_uploads"
    ) -> str:
        """
        Save markdown content to a file with proper naming convention.

        Args:
            markdown_content: The markdown content to save
            template_name: Name of the template
            username: Username of the user uploading
            base_dir: Base directory for uploads (default: document_uploads)

        Returns:
            Path to the saved markdown file
        """
        # Create document_uploads directory at root level if it doesn't exist
        root_dir = Path(__file__).parent.parent.parent.parent  # Go up to project root
        upload_dir = root_dir / base_dir
        upload_dir.mkdir(parents=True, exist_ok=True)

        # Generate filename: templatename_username_timestamp.md
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        # Clean template name for filename (remove special characters)
        clean_template_name = "".join(c for c in template_name if c.isalnum() or c in (' ', '-', '_')).strip()
        clean_template_name = clean_template_name.replace(' ', '_')
        filename = f"{clean_template_name}_{username}_{timestamp}.md"

        file_path = upload_dir / filename

        # Save the markdown content
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(markdown_content)

        return str(file_path)
    
    @staticmethod
    def ocr_pdf_with_openai(file_path: str) -> Optional[str]:
        """
        Use OpenAI Vision API to OCR an image-based PDF.
        Converts PDF pages to images and sends them to GPT-4 Vision.
        
        Args:
            file_path: Path to the PDF file
            
        Returns:
            Extracted text as markdown, or None if OCR fails
        """
        if not OPENAI_AVAILABLE:
            logging.error("OpenAI package not installed")
            return None
        
        if not PDF2IMAGE_AVAILABLE:
            logging.error("pdf2image package not installed")
            return None
        
        from ..config import settings
        
        if not settings.openai_api_key:
            logging.error("OpenAI API key not configured")
            return None
        
        try:
            # Convert PDF pages to images
            logging.info(f"Converting PDF to images: {file_path}")
            images = convert_from_path(file_path, dpi=150)
            logging.info(f"Converted {len(images)} pages to images")
            
            client = OpenAI(api_key=settings.openai_api_key)
            all_text = []
            
            for i, image in enumerate(images):
                logging.info(f"Processing page {i+1} with OpenAI Vision")
                
                # Convert PIL image to base64
                buffered = BytesIO()
                image.save(buffered, format="PNG")
                img_base64 = base64.b64encode(buffered.getvalue()).decode('utf-8')
                
                # Call OpenAI Vision API
                response = client.chat.completions.create(
                    model="gpt-4o",
                    messages=[
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "text",
                                    "text": "Please extract all the text from this document image. Preserve the structure and formatting as much as possible. Output the text in markdown format. If there are form fields or placeholders, preserve them. Do not add any commentary, just output the extracted text."
                                },
                                {
                                    "type": "image_url",
                                    "image_url": {
                                        "url": f"data:image/png;base64,{img_base64}"
                                    }
                                }
                            ]
                        }
                    ],
                    max_tokens=4096
                )
                
                page_text = response.choices[0].message.content
                if page_text:
                    all_text.append(f"<!-- Page {i+1} -->\n{page_text}")
                    logging.info(f"Page {i+1} extracted {len(page_text)} characters")
            
            result = "\n\n".join(all_text)
            logging.info(f"Total OCR result: {len(result)} characters")
            return result
            
        except Exception as e:
            logging.error(f"OpenAI Vision OCR failed: {e}")
            return None
    
    @staticmethod
    def ocr_image_with_openai(file_path: str) -> Optional[str]:
        """
        Use OpenAI Vision API to OCR an image file.
        
        Args:
            file_path: Path to the image file
            
        Returns:
            Extracted text as markdown, or None if OCR fails
        """
        if not OPENAI_AVAILABLE:
            logging.error("OpenAI package not installed")
            return None
        
        if not PIL_AVAILABLE:
            logging.error("PIL/Pillow package not installed")
            return None
        
        from ..config import settings
        
        if not settings.openai_api_key:
            logging.error("OpenAI API key not configured")
            return None
        
        try:
            logging.info(f"Processing image with OpenAI Vision: {file_path}")
            
            # Read and encode the image
            with open(file_path, 'rb') as f:
                img_base64 = base64.b64encode(f.read()).decode('utf-8')
            
            # Determine image type from extension
            ext = os.path.splitext(file_path)[1].lower()
            mime_type = {
                '.jpg': 'image/jpeg',
                '.jpeg': 'image/jpeg',
                '.png': 'image/png',
                '.tiff': 'image/tiff',
                '.tif': 'image/tiff',
                '.bmp': 'image/bmp'
            }.get(ext, 'image/png')
            
            client = OpenAI(api_key=settings.openai_api_key)
            
            response = client.chat.completions.create(
                model="gpt-4o",
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": "Please extract all the text from this document image. Preserve the structure and formatting as much as possible. Output the text in markdown format. If there are form fields or placeholders, preserve them. Do not add any commentary, just output the extracted text."
                            },
                            {
                                "type": "image_url",
                                "image_url": {
                                    "url": f"data:{mime_type};base64,{img_base64}"
                                }
                            }
                        ]
                    }
                ],
                max_tokens=4096
            )
            
            result = response.choices[0].message.content
            logging.info(f"Image OCR result: {len(result) if result else 0} characters")
            return result
            
        except Exception as e:
            logging.error(f"OpenAI Vision OCR failed for image: {e}")
            return None
